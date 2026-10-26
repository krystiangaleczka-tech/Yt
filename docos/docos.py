# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl>=3.1"]
# ///
"""DocOS — system scenopisarski dla faceless dokumentów YouTube.

Źródłem prawdy jest jeden plik projektu (JSON). Arkusz .xlsx to widok do planowania
(heatmapa energii, listy rozwijane), który można edytować i wczytać z powrotem.

Workflow z Claude (Claude sam jest "silnikiem" — bez kluczy API):
    uv run docos.py new projekt.json "Tytuł"
    uv run docos.py system projekt.json          # globalne DNA odcinka
    uv run docos.py next projekt.json            # prompt dla kolejnego pustego beatu
    uv run docos.py save projekt.json B03 wynik.md   # zapis wyniku, auto N-1 dla kolejnego
    uv run docos.py check projekt.json           # krzywa energii, fakty, czasy
    uv run docos.py export projekt.json          # .xlsx + script.md + narration.txt + render_plan.json
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

# ───────────────────────── słowniki systemu ─────────────────────────

SEGMENTS = {
    # segment: (domyślna dynamika, zakres średniej energii)
    "COLD_OPEN": ("ACTION", (8, 10)),
    "CONTEXT": ("STORY", (4, 7)),
    "INVESTIGATION": ("STORY", (4, 7)),
    "TWIST": ("EUREKA", (7, 9)),
    "RESOLUTION": ("EUREKA", (6, 9)),
    "OUTRO": ("CHILL", (2, 6)),
}

# Looped Arc dla segmentu INVESTIGATION + klasyczne typologie
TYPOLOGIES = [
    "INTRODUCTION", "RISING_ACTION", "CONFLICT", "COMEBACK", "EVIDENCE",
    "CLIMAX", "FALLING_ACTION", "PAYOFF", "TRANSITION", "OPEN_QUESTION",
]

DYNAMICS = {
    # profil: (energia, min ujęć, max ujęć, min s/ujęcie, max s/ujęcie, opis dla AI) — tempo dokumentu: słucha się, nie skacze
    "ACTION": (10, 4, 7, 1.5, 3.5,
               "Szybsza seria: ujęcia 1.5–3.5 s, krótkie zdania. Tylko hooki i kulminacje, nie częściej niż co kilka beatów."),
    "STORY": (5, 2, 5, 4.5, 7.0,
              "Rytm dokumentu: ujęcia 4.5–7 s z powolnym najazdem, czas na obejrzenie zdjęcia."),
    "CHILL": (2, 1, 3, 6.0, 12.0,
              "Oddech: długie ujęcia 6–12 s, cisza, refleksja."),
    "EUREKA": (8, 3, 6, 1.5, 7.0,
               "Hybrid: 2–3 krótsze cięcia (szukanie), potem długie ujęcie z najazdem (odkrycie)."),
    "STATIC": (3, 1, 2, 6.0, 14.0,
               "Dokument: jedno długie, stabilne ujęcie, czas na przeczytanie."),
}

CHAOS = {
    "NONE": "Trzymaj się planu z pola Akcja.",
    "MICRO": "Wprowadź drobną anomalię lub detal ze źródeł (np. brakująca strona, dziwna adnotacja). Nie zmienia celu beatu.",
    "MINI": "Komplikacja śledztwa: nowa informacja z rejestru faktów podważa to, co widz założył. Rozwiąż ją lub zostaw jako otwartą pętlę.",
    "MACRO": "Odwrócenie tezy rozdziału: fakt z rejestru wywraca dotychczasową interpretację. Kolejne beaty muszą skorygować kurs.",
}

FACT_STATUS = ["VERIFIED", "DISPUTED", "UNVERIFIED", "DRAMATIZATION"]

ASSETS = {
    "ARCHIWUM": "prawdziwy materiał archiwalny / stock / domena publiczna (podaj czego szukać)",
    "AI_OBRAZ": "kadr generowany AI (podaj prompt po angielsku w stylu globalnym)",
    "DOKUMENT": "skan / zdjęcie dokumentu, raportu, wycinka (najazd Ken Burns)",
    "TYTUŁ": "napis maszynopisem: data, miejsce, etykieta",
    "PORÓWNANIE": "plansza dwóch kart: LEWA vs PRAWA z podpisami",
    "MAPA": "mapa z zaznaczeniem miejsca / trasy",
    "OŚ_CZASU": "oś czasu z wydarzeniami",
    "CYTAT": "cytat na planszy (tylko dosłowny, ze źródła)",
    "PRASA": 'plansza prasowa: dosłowny cytat z artykułu + "— Tytuł pisma, data" (własny layout, nie podróbka gazety)',
    "PODPIS": "belka z imieniem i funkcją osoby (nie sugeruj, że osoba jest na zdjęciu, jeśli nie jest)",
    "WAVEFORM": "żywy oscylogram z PRAWDZIWEGO nagrania (FILE: plik.wav)",
    "SPEKTRO": "przewijany spektrogram liczony z PRAWDZIWEGO nagrania (FILE: plik.wav)",
    "LUPA": 'przybliżenie detalu zdjęcia z ramką (FILE: plik ZOOM: x,y,w jako ułamki) + "podpis detalu"',
    "LICZNIK": 'animowany licznik liczby z faktów: "0 → 3,219 km"',
    "DIAGRAM": "diagram objaśniający (np. SOFAR: przekrój oceanu, droga dźwięku)",
    "SPLIT": 'dzielony ekran dwóch zdjęć (FILES: a | b) z podpisami "lewy" / "prawy"',
    "STOS": "stos zdjęć spadających na stół (FILES: a | b | c)",
    "QA": 'rozmowa: Q: "pytanie narratora" A: "dosłowny cytat" — źródło (tylko prawdziwe cytaty)',
    "HASŁO": 'kinetyczne hasło: słowa wskakują jedno po drugim, "TEKST"',
}

MAX_RUN = {"ACTION": 3, "LOW": 2}  # max ACTION pod rząd; max CHILL/STATIC pod rząd
LOW = {"CHILL", "STATIC"}

# ───────────────────────── model projektu ─────────────────────────

TEMPLATE_META = {
    "title": "",
    "logline": "",
    "thesis": "",
    "curiosity_gap": "",
    "narrator": "Spokojny, rzeczowy narrator dokumentalny w 3. osobie. Niski ton, krótkie zdania, zero sensacji bez dowodu.",
    "audience_relation": "Prowadzi widza przez akta jak śledczy. Czasem zwraca się na 'ty' w kluczowych momentach.",
    "lexicon": "",
    "bans": "Zero lania wody. Żadnych zmyślonych nazwisk, dat, liczb i cytatów. Nie udawaj, że rekonstrukcja to archiwum.",
    "language": "English",
    "target_minutes": 15,
    "wpm": 140,
    "visual_style": "1980s–90s archival documentary look, black and white, film grain, vignette, typewriter titles, red/black comparison cards",
    "master_outline": "",
}


def load(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def dump(path: Path, prj: dict) -> None:
    path.write_text(json.dumps(prj, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def new_project(title: str) -> dict:
    meta = dict(TEMPLATE_META, title=title)
    return {"meta": meta, "chapters": [], "subchapters": [], "facts": [], "beats": []}


def beat_index(prj: dict, beat_id: str) -> int:
    for i, b in enumerate(prj["beats"]):
        if b["id"].upper() == beat_id.upper():
            return i
    sys.exit(f"Nie ma beatu {beat_id}")


def by_id(items: list[dict], key: str) -> dict:
    return next((x for x in items if str(x["id"]) == str(key)), {})


def effective_dynamic(prj: dict, b: dict) -> str:
    """Smart setup: beat > rozdział > segment ("Rozdział to pogoda, beat to piorun")."""
    if b.get("dynamic"):
        return b["dynamic"]
    ch = by_id(prj["chapters"], b.get("chapter", ""))
    if ch.get("default_dynamic"):
        return ch["default_dynamic"]
    return SEGMENTS.get(b.get("segment", ""), ("STORY",))[0]


def shot_target(dyn: str, seconds: float) -> tuple[int, int, int]:
    e, lo, hi, smin, smax = DYNAMICS[dyn][:5]
    mean = (smin + smax) / 2
    n = max(lo, min(hi, round(seconds / mean))) if seconds else lo
    return n, lo, hi


def words_for(prj: dict, seconds: float) -> int:
    return round(seconds * prj["meta"].get("wpm", 140) / 60)


def summary_of(prj: dict, i: int) -> str:
    if i < 0:
        return "(początek filmu)"
    b = prj["beats"][i]
    return b.get("link") or f"(beat {b['id']} jeszcze nienapisany — plan: {b.get('action', '')})"


# ───────────────────────── prompty ─────────────────────────

def system_prompt(prj: dict) -> str:
    m = prj["meta"]
    chapters = "\n".join(
        f"  Rozdział {c['id']}: {c['title']} — {c.get('description', '')} [segment {c.get('segment', '?')}]"
        for c in prj["chapters"]
    ) or "  (brak rozdziałów)"
    facts = "\n".join(
        f"  {f['id']} [{f['status']}] {f['claim']} — źródło: {f.get('source', '?')}" for f in prj["facts"]
    ) or "  (pusty — najpierw research!)"
    return f"""### USTAWIENIA GLOBALNE ODCINKA (System Prompt — DocOS)
Piszemy faceless dokument YouTube MODUŁOWO: beat po beacie. Nie widzisz całego filmu, tylko bieżący beat
i streszczenia N-1/N-2. Każdy beat to 1 emocja wejściowa → 1 emocja wyjściowa.

[TYTUŁ]: {m['title']}
[LOGLINE]: {m['logline']}
[TEZA / KĄT]: {m['thesis']}
[CURIOSITY GAP]: {m['curiosity_gap']}
[NARRATOR]: {m['narrator']}
[RELACJA Z WIDZEM]: {m['audience_relation']}
[SŁOWNIK]: {m['lexicon'] or '-'}
[ZAKAZY]: {m['bans']}
[JĘZYK NARRACJI]: {m['language']}
[STYL WIZUALNY]: {m['visual_style']}
[MASTER OUTLINE]: {m['master_outline'] or '-'}
[ROZDZIAŁY]:
{chapters}

### ZASADA PRAWDY (dokument, nie fikcja)
Jedyne dozwolone twarde dane (liczby, daty, nazwiska, cytaty) pochodzą z REJESTRU FAKTÓW:
{facts}
- VERIFIED: podawaj wprost.
- DISPUTED: tylko z zastrzeżeniem ("według...", "do dziś nie wiadomo").
- UNVERIFIED: nie stwierdzaj jako faktu; możesz postawić jako pytanie.
- DRAMATIZATION: wolno pokazać jako rekonstrukcję; w shot liście oznacz ujęcie jako (REKONSTRUKCJA).
Jeśli beat wymaga faktu, którego nie ma w rejestrze — NIE wymyślaj go. Wstaw [POTRZEBNY RESEARCH: czego brakuje].

### ZASADA ZWROTU (kolumna Chaos)
Chaos to licencja na złamanie planu beatu, ale w dokumencie zwrot musi wynikać z faktów:
""" + "\n".join(f"- {k}: {v}" for k, v in CHAOS.items()) + """
Konsekwencję zwrotu opisz w 'Połączenie do N-1', żeby kolejne beaty skorygowały kurs.

### PROFILE DYNAMIKI WIZUALNEJ
""" + "\n".join(f"- {k} (energia {v[0]}, {v[1]}–{v[2]} ujęć, {v[3]}–{v[4]} s/ujęcie): {v[5]}" for k, v in DYNAMICS.items()) + """

### TYPY MATERIAŁU W SHOT LIŚCIE
""" + "\n".join(f"- {k}: {v}" for k, v in ASSETS.items()) + """
Shot list jest produkcyjny, nie literacki: kąt kamery, ruch, detal, czas (Wide, ECU, Macro, POV, Match Cut, Ken Burns, Frame-in-frame).
Bohater pod presją czasu działa — nie czeka biernie. Logika akcji ponad efekt.
"""


def beat_prompt(prj: dict, i: int) -> str:
    b = prj["beats"][i]
    m = prj["meta"]
    dyn = effective_dynamic(prj, b)
    sec = float(b.get("seconds") or 30)
    n, lo, hi = shot_target(dyn, sec)
    d = DYNAMICS[dyn]
    ch = by_id(prj["chapters"], b.get("chapter", ""))
    sub = by_id(prj["subchapters"], b.get("subchapter", ""))
    chaos = b.get("chaos") or "NONE"
    fact_ids = b.get("facts") or []
    facts = [by_id(prj["facts"], fid) for fid in fact_ids]
    fact_lines = "\n".join(
        f"- {f['id']} [{f['status']}] {f['claim']} (źródło: {f.get('source', '?')})" for f in facts if f
    ) or "- (brak przypisanych faktów — nie podawaj żadnych twardych danych)"
    return f"""### ROLA: Scenarzysta i reżyser faceless dokumentu YouTube. Styl: surowy, filmowy, rzeczowy. ZERO lania wody.
[TYTUŁ]: {m['title']} | [TEZA]: {m['thesis']}
[SEGMENT]: {b.get('segment', '')}
[ROZDZIAŁ]: {ch.get('id', '-')} ({ch.get('title', '-')}: {ch.get('description', '')})
[PODROZDZIAŁ]: {sub.get('id', '-')} ({sub.get('title', '-')}: {sub.get('description', '')})
[SCENA]: {b.get('scene_no', '')} ({b.get('scene_desc', '')})
[BEAT]: {b['id']} | [TYPOLOGIA]: {b.get('typology', '')} | [DYNAMIKA]: {dyn}
[KONTEKST N-2]: {summary_of(prj, i - 2)}
[KONTEKST N-1]: {summary_of(prj, i - 1)}

### ZADANIE: Napisz BEAT skryptu.
Akcja beatu: '{b.get('action', '')}'
Cel logiczny (co widz ma zrozumieć): '{b.get('goal', '')}'
Zmiana emocjonalna (Turn): '{b.get('turn', '')}'
Długość: ~{sec:.0f} s narracji = ok. {words_for(prj, sec)} słów (±15%). Język narracji: {m['language']}.

### FAKTY DLA TEGO BEATU (jedyne dozwolone twarde dane):
{fact_lines}

### PROCEDURA ZWROTU (Chaos: {chaos}):
{CHAOS.get(chaos, CHAOS['NONE'])}

### WYMAGANIA:
1. NARRACJA: Show, don't tell. Rytm zdań dopasuj do dynamiki {dyn}: {d[5]}
2. SHOT LIST: {n} ujęć (dozwolony zakres {lo}–{hi}), po {d[3]}–{d[4]} s. Suma czasów ≈ {sec:.0f} s.
   Każde ujęcie w formacie: [WIZ n] (typ kadru / ruch) TYP_MATERIAŁU: opis (X s)
   TYP_MATERIAŁU ∈ {', '.join(ASSETS)}. Dla AI_OBRAZ dopisz na końcu opisu: PROMPT: "<prompt EN w stylu: {m['visual_style']}>"

### WYMAGANY FORMAT WYJŚCIOWY (dokładnie te nagłówki):
Narracja:
<tekst lektora>

Shot List:
[WIZ 1] ...

Fakty użyte: <ID faktów, np. F1, F3 — albo 'brak'>
Połączenie do N-1: <jedno zdanie streszczające beat dla kolejnego wiersza>
"""


# ───────────────────────── parsowanie wyniku ─────────────────────────

SHOT_RE = re.compile(
    r"\[WIZ\s*(\d+)\]\s*(?:\(([^)]*)\))?\s*([A-ZĄĆĘŁŃÓŚŹŻ_]+)?\s*:?\s*(.*?)\s*\((\d+(?:[.,]\d+)?)\s*s\)\s*$",
    re.IGNORECASE,
)


def parse_result(text: str) -> dict:
    def section(name: str, nxt: list[str]) -> str:
        pat = rf"{name}\s*:?\s*\n?(.*?)(?=^\s*(?:{'|'.join(nxt)})\s*:|\Z)"
        mm = re.search(pat, text, re.S | re.M | re.I)
        return mm.group(1).strip() if mm else ""

    narration = section("Narracja", ["Shot List", "Scenopis", "Fakty użyte", "Połączenie"])
    shots_raw = section(r"(?:Shot List|Scenopis Obrazkowy(?: \(Shot List\))?)", ["Fakty użyte", "Połączenie"])
    facts_used = section("Fakty użyte", ["Połączenie"])
    link = section("Połączenie do N-1", ["ZZZ"])
    shots = []
    for line in shots_raw.splitlines():
        line = line.strip().lstrip("-• ")
        if not line.upper().startswith("[WIZ"):
            continue
        mm = SHOT_RE.match(line)
        if mm:
            num, framing, asset, desc, dur = mm.groups()
            prompt = ""
            pm = re.search(r'PROMPT:\s*"?(.*?)"?\s*$', desc)
            if pm:
                prompt, desc = pm.group(1), desc[: pm.start()].rstrip(" .;")
            shots.append({"n": int(num), "framing": framing or "", "asset": (asset or "").upper(),
                          "desc": desc, "prompt": prompt, "dur": float(dur.replace(",", "."))})
        else:
            shots.append({"n": len(shots) + 1, "framing": "", "asset": "", "desc": line, "prompt": "", "dur": 0.0})
    ids = re.findall(r"F\d+", facts_used.upper())
    return {"narration": narration, "shots": shots, "facts_used": ids, "link": link.splitlines()[0] if link else ""}


# ───────────────────────── walidacja ─────────────────────────

def turn_polarity(turn: str) -> str:
    end = turn.split("→")[-1] if "→" in turn else turn.split("->")[-1]
    if "[+" in end:
        return "+"
    if "[-" in end or "[—" in end:
        return "-"
    return "?"


def check(prj: dict) -> list[tuple[str, str]]:
    out: list[tuple[str, str]] = []
    beats = prj["beats"]
    facts = {f["id"]: f for f in prj["facts"]}
    dyns = [effective_dynamic(prj, b) for b in beats]

    # zasada 3 beatów
    run, prev = 0, None
    for b, d in zip(beats, dyns):
        key = "LOW" if d in LOW else d
        run = run + 1 if key == prev else 1
        prev = key
        if key in MAX_RUN and run > MAX_RUN[key]:
            out.append(("ENERGIA", f"{b['id']}: {run}× {'CHILL/STATIC' if key == 'LOW' else key} pod rząd (max {MAX_RUN[key]}) — zmień dynamikę."))

    # segment jako master fader
    seg_e: dict[str, list[int]] = {}
    for b, d in zip(beats, dyns):
        seg_e.setdefault(b.get("segment", "?"), []).append(DYNAMICS[d][0])
    for seg, es in seg_e.items():
        if seg in SEGMENTS:
            lo, hi = SEGMENTS[seg][1]
            avg = sum(es) / len(es)
            if not lo <= avg <= hi:
                out.append(("ENERGIA", f"Segment {seg}: średnia energia {avg:.1f}, oczekiwane {lo}–{hi}."))

    # turny
    pol_run, last = 0, None
    for b in beats:
        t = b.get("turn", "")
        arrows = t.count("→") + t.count("->")
        if not t:
            out.append(("TURN", f"{b['id']}: brak Turnu — jeśli nic się nie zmienia, beat jest zbędny."))
        elif arrows > 1:
            out.append(("TURN", f"{b['id']}: wielostopniowy Turn ({t}). Zasada: 1 emocja wejściowa → 1 wyjściowa."))
        p = turn_polarity(t)
        pol_run = pol_run + 1 if p == last and p != "?" else 1
        last = p
        if pol_run == 4:
            out.append(("TURN", f"{b['id']}: 4 beaty z rzędu kończą się na [{p}] — daj kontrapunkt."))

    # fakty
    for b in beats:
        for fid in b.get("facts", []):
            f = facts.get(fid)
            if not f:
                out.append(("FAKTY", f"{b['id']}: nieznany fakt {fid}."))
            elif f["status"] == "UNVERIFIED":
                out.append(("FAKTY", f"{b['id']}: używa {fid} (UNVERIFIED) — zweryfikuj albo sformułuj jako pytanie."))
        nar = b.get("narration", "")
        if nar:
            allowed = " ".join(facts[f]["claim"] for f in b.get("facts", []) if f in facts)
            allowed_nums = set(re.findall(r"\d[\d.,]*", allowed))
            for num in set(re.findall(r"\d[\d.,]*\d|\d", nar)):
                if num.strip(".,") not in {a.strip(".,") for a in allowed_nums}:
                    out.append(("FAKTY", f"{b['id']}: liczba '{num}' w narracji nie występuje w przypisanych faktach."))
            if "[POTRZEBNY RESEARCH" in nar.upper():
                out.append(("FAKTY", f"{b['id']}: narracja zgłasza brak researchu."))

    # czasy i ujęcia
    total = 0.0
    for b, d in zip(beats, dyns):
        sec = float(b.get("seconds") or 0)
        total += sec
        nar = b.get("narration", "")
        if nar and sec:
            want = words_for(prj, sec)
            got = len(nar.split())
            if abs(got - want) > 0.25 * want:
                out.append(("CZAS", f"{b['id']}: {got} słów, budżet ~{want} ({sec:.0f} s)."))
        shots = b.get("shots", [])
        if shots:
            _, lo, hi = shot_target(d, sec)
            if not lo <= len(shots) <= hi:
                out.append(("UJĘCIA", f"{b['id']}: {len(shots)} ujęć, profil {d} wymaga {lo}–{hi}."))
            s = sum(x["dur"] for x in shots)
            if sec and abs(s - sec) > 0.3 * sec:
                out.append(("UJĘCIA", f"{b['id']}: suma ujęć {s:.1f} s vs beat {sec:.0f} s."))
            bad = [x["n"] for x in shots if x["asset"] and x["asset"] not in ASSETS]
            if bad:
                out.append(("UJĘCIA", f"{b['id']}: nieznany typ materiału w WIZ {bad}."))
    # różnorodność obrazu: unikalne materiały, powtórzenia, maszynopis, monotonia typów
    all_shots = [(b["id"], s) for b in beats for s in b.get("shots", [])]
    if all_shots:
        def key(s):
            m = re.search(r"FILES?:\s*([^\s(]+)", s["desc"])
            return m.group(1) if m else s["asset"] + ":" + s["desc"][:30]
        uses: dict[str, list[str]] = {}
        for bid, s in all_shots:
            uses.setdefault(key(s), []).append(bid)
        want = prj["meta"].get("min_unique_visuals", 20)
        if len(uses) < want:
            out.append(("RÓŻNORODNOŚĆ", f"Tylko {len(uses)} różnych materiałów (min {want})."))
        for k_, bs in uses.items():
            if len(bs) > 2:
                out.append(("RÓŻNORODNOŚĆ", f"'{k_[:40]}' użyty {len(bs)}× ({', '.join(bs)}) — max 2."))
        typed = sum(1 for _, s in all_shots if s["asset"] == "TYTUŁ" and "stamp" not in s["desc"])
        cap = max(3, round(len(all_shots) * 0.06))
        if typed > cap:
            out.append(("RÓŻNORODNOŚĆ", f"{typed} napisów maszynopisem (max {cap}) — użyj HASŁO, LICZNIK, LUPA, WAVEFORM."))
        run, last = 0, None
        for bid, s in all_shots:
            run = run + 1 if s["asset"] == last else 1
            last = s["asset"]
            if run == 3:
                out.append(("RÓŻNORODNOŚĆ", f"{bid}: 3× {s['asset']} pod rząd."))
        text_assets = {"TYTUŁ", "HASŁO", "LICZNIK", "CYTAT", "PRASA", "QA", "OŚ_CZASU"}
        for (b1, s1), (b2, s2) in zip(all_shots, all_shots[1:]):
            if s1["asset"] in text_assets and s2["asset"] in text_assets:
                out.append(("RÓŻNORODNOŚĆ", f"{b2}: dwie plansze tekstowe obok siebie ({s1['asset']} → {s2['asset']}) — wstaw obraz pomiędzy."))
            st1, st2 = s1.get("style"), s2.get("style")
            if st1 and st1 == st2 and st1 != "auto":
                out.append(("RÓŻNORODNOŚĆ", f"{b2}: ten sam styl '{st1}' dwa razy z rzędu."))
        trs = [s.get("transition") for _, s in all_shots]
        for i in range(2, len(trs)):
            if trs[i] and trs[i] == trs[i - 1] == trs[i - 2] and trs[i] != "cut":
                out.append(("RÓŻNORODNOŚĆ", f"{all_shots[i][0]}: przejście '{trs[i]}' 3× z rzędu."))

    target = prj["meta"].get("target_minutes", 0) * 60
    if target and beats and abs(total - target) > 0.15 * target:
        out.append(("CZAS", f"Suma beatów {total / 60:.1f} min vs cel {target / 60:.0f} min."))

    # zszywacze
    for i, b in enumerate(beats[:-1]):
        if b.get("narration") and not b.get("link"):
            out.append(("CIĄGŁOŚĆ", f"{b['id']}: brak 'Połączenia do N-1' — kolejny beat straci kontekst."))
    return out


# ───────────────────────── eksporty ─────────────────────────

FILL = {"ACTION": "E74C3C", "STORY": "F5B041", "CHILL": "5DADE2", "EUREKA": "AF7AC5", "STATIC": "48C9B0"}
BEAT_COLS = [
    ("id", "Beat ID"), ("segment", "Segment"), ("scene_no", "Scena Nr"), ("scene_desc", "Scena Opis"),
    ("chapter", "Rozdział Nr"), ("subchapter", "Podrozdział Nr"), ("action", "Beat Akcja"),
    ("goal", "Cel Beatu"), ("turn", "Turn (Zmiana)"), ("typology", "Typologia"), ("dynamic", "Dynamika Wizualna"),
    ("chaos", "Chaos"), ("facts", "Fakty (ID)"), ("seconds", "Czas (s)"),
]


def export_xlsx(prj: dict, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    ws = wb.active
    ws.title = "Ustawienia"
    ws.append(["Parametr", "Wartość"])
    for k, v in prj["meta"].items():
        ws.append([k, v])
    ws.append(["SYSTEM_PROMPT (tylko odczyt)", system_prompt(prj)])
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows(min_row=2):
        row[1].alignment = wrap

    wc = wb.create_sheet("Rozdziały")
    wc.append(["Typ", "ID", "Rozdział Nr (dla podrozdziału)", "Tytuł", "Opis", "Segment", "Domyślna dynamika"])
    for c in prj["chapters"]:
        wc.append(["ROZDZIAŁ", c["id"], "", c["title"], c.get("description", ""), c.get("segment", ""), c.get("default_dynamic", "")])
    for s in prj["subchapters"]:
        wc.append(["PODROZDZIAŁ", s["id"], s.get("chapter", ""), s["title"], s.get("description", ""), "", ""])

    wf = wb.create_sheet("Fakty")
    wf.append(["ID", "Twierdzenie", "Status", "Źródło", "URL", "Notatka"])
    for f in prj["facts"]:
        wf.append([f["id"], f["claim"], f["status"], f.get("source", ""), f.get("url", ""), f.get("note", "")])
    dv = DataValidation(type="list", formula1='"' + ",".join(FACT_STATUS) + '"', allow_blank=True)
    wf.add_data_validation(dv)
    dv.add("C2:C500")
    for col, w in zip("ABCDEF", (6, 70, 15, 35, 40, 30)):
        wf.column_dimensions[col].width = w

    wbt = wb.create_sheet("Beaty")
    heads = [h for _, h in BEAT_COLS] + ["Dynamika (efektywna)", "Energia", "Kontekst N-2", "Kontekst N-1",
                                         "ZŁOTA FORMUŁA (PROMPT)", "Narracja", "Połączenie do N-1", "Timestamp"]
    wbt.append(heads)
    t = 0.0
    for i, b in enumerate(prj["beats"]):
        d = effective_dynamic(prj, b)
        row = [", ".join(b.get(k, [])) if k == "facts" else b.get(k, "") for k, _ in BEAT_COLS]
        row += [d, DYNAMICS[d][0], summary_of(prj, i - 2), summary_of(prj, i - 1), beat_prompt(prj, i),
                b.get("narration", ""), b.get("link", ""), f"{int(t // 60)}:{int(t % 60):02d}"]
        t += float(b.get("seconds") or 0)
        wbt.append(row)
    n = len(prj["beats"]) + 1
    lists = {"Segment": list(SEGMENTS), "Typologia": TYPOLOGIES, "Dynamika Wizualna": list(DYNAMICS), "Chaos": list(CHAOS)}
    for idx, (_, h) in enumerate(BEAT_COLS, start=1):
        if h in lists:
            col = wbt.cell(row=1, column=idx).column_letter
            v = DataValidation(type="list", formula1='"' + ",".join(lists[h]) + '"', allow_blank=True)
            wbt.add_data_validation(v)
            v.add(f"{col}2:{col}500")
    eff_col = wbt.cell(row=1, column=len(BEAT_COLS) + 1).column_letter
    dyn_col = wbt.cell(row=1, column=[h for _, h in BEAT_COLS].index("Dynamika Wizualna") + 1).column_letter
    for col in (dyn_col, eff_col):
        for name, color in FILL.items():
            wbt.conditional_formatting.add(
                f"{col}2:{col}500",
                CellIsRule(operator="equal", formula=[f'"{name}"'], fill=PatternFill("solid", fgColor=color)))
    for c in wbt[1]:
        c.font = bold
    for row in wbt.iter_rows(min_row=2):
        for c in row:
            c.alignment = wrap
    widths = [8, 14, 8, 24, 8, 10, 40, 32, 26, 16, 14, 9, 12, 8, 14, 8, 30, 30, 60, 60, 40, 9]
    for i, w in enumerate(widths, start=1):
        wbt.column_dimensions[wbt.cell(row=1, column=i).column_letter].width = w
    wbt.freeze_panes = "B2"

    ws2 = wb.create_sheet("Shot List")
    ws2.append(["Beat", "WIZ", "Kadr / ruch", "Materiał", "Opis", "Prompt AI", "Czas (s)"])
    for b in prj["beats"]:
        for s in b.get("shots", []):
            ws2.append([b["id"], s["n"], s["framing"], s["asset"], s["desc"], s["prompt"], s["dur"]])
    for col, w in zip("ABCDEFG", (8, 5, 24, 13, 70, 60, 8)):
        ws2.column_dimensions[col].width = w

    we = wb.create_sheet("Energia")
    we.append(["Beat", "Energia", "Dynamika"])
    for b in prj["beats"]:
        d = effective_dynamic(prj, b)
        we.append([b["id"], DYNAMICS[d][0], d])
    if prj["beats"]:
        ch = LineChart()
        ch.title = "Krzywa energii (cel: fala 3–10)"
        ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, 10
        ch.add_data(Reference(we, min_col=2, min_row=1, max_row=n), titles_from_data=True)
        ch.set_categories(Reference(we, min_col=1, min_row=2, max_row=n))
        ch.width, ch.height = 26, 9
        we.add_chart(ch, "E2")
    issues = check(prj)
    wk = wb.create_sheet("Kontrola")
    wk.append(["Kategoria", "Problem"])
    for cat, msg in issues or [("OK", "Brak problemów.")]:
        wk.append([cat, msg])
    wk.column_dimensions["B"].width = 120
    wb.save(path)


def import_xlsx(prj: dict, path: Path) -> dict:
    """Wczytuje z arkusza planowanie (Ustawienia, Rozdziały, Fakty, kolumny planu Beatów). Wyników AI nie nadpisuje."""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    for k, v in wb["Ustawienia"].iter_rows(min_row=2, values_only=True):
        if k in prj["meta"]:
            prj["meta"][k] = v if v is not None else ""
    prj["chapters"], prj["subchapters"] = [], []
    for typ, cid, parent, title, desc, seg, dyn in wb["Rozdziały"].iter_rows(min_row=2, values_only=True):
        if not cid:
            continue
        if typ == "PODROZDZIAŁ":
            prj["subchapters"].append({"id": str(cid), "chapter": str(parent or ""), "title": title or "", "description": desc or ""})
        else:
            prj["chapters"].append({"id": str(cid), "title": title or "", "description": desc or "",
                                    "segment": seg or "", "default_dynamic": dyn or ""})
    prj["facts"] = [
        {"id": r[0], "claim": r[1] or "", "status": r[2] or "UNVERIFIED", "source": r[3] or "", "url": r[4] or "", "note": r[5] or ""}
        for r in wb["Fakty"].iter_rows(min_row=2, values_only=True) if r[0]
    ]
    old = {b["id"]: b for b in prj["beats"]}
    beats = []
    for r in wb["Beaty"].iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        b = dict(old.get(r[0], {}))
        for (k, _), v in zip(BEAT_COLS, r):
            if k == "facts":
                b[k] = [x.strip() for x in str(v or "").split(",") if x.strip()]
            elif k == "seconds":
                b[k] = float(v or 0)
            else:
                b[k] = "" if v is None else str(v)
        beats.append(b)
    prj["beats"] = beats
    return prj


def export_script(prj: dict) -> str:
    out = [f"# {prj['meta']['title']}\n", f"_{prj['meta']['logline']}_\n"]
    t = 0.0
    for b in prj["beats"]:
        d = effective_dynamic(prj, b)
        out.append(f"\n## {b['id']} · {int(t // 60)}:{int(t % 60):02d} · {b.get('segment', '')} · {d} · Turn: {b.get('turn', '')}\n")
        out.append(b.get("narration") or f"> (do napisania) {b.get('action', '')}")
        if b.get("shots"):
            out.append("\n**Shot list**\n")
            out += [f"- WIZ {s['n']} ({s['framing']}) **{s['asset']}**: {s['desc']} — {s['dur']} s" for s in b["shots"]]
        t += float(b.get("seconds") or 0)
    used = sorted({f for b in prj["beats"] for f in b.get("facts_used", b.get("facts", []))})
    out.append("\n## Źródła (do opisu filmu)\n")
    for f in prj["facts"]:
        if f["id"] in used and f["status"] != "DRAMATIZATION":
            out.append(f"- {f['claim']} — {f.get('source', '')} {f.get('url', '')}".rstrip())
    if any(f["status"] == "DRAMATIZATION" for f in prj["facts"] if f["id"] in used):
        out.append("\n_Film zawiera rekonstrukcje i kadry wygenerowane przez AI._")
    return "\n".join(out) + "\n"


def render_plan(prj: dict) -> dict:
    """Plan dla pipeline'u wideo (pilot/): lektor per beat + lista ujęć z promptami AI i czasem startu."""
    plan, t = [], 0.0
    for b in prj["beats"]:
        st = t
        for s in b.get("shots", []):
            plan.append({"beat": b["id"], "wiz": s["n"], "start": round(st, 2), "dur": s["dur"], "asset": s["asset"],
                         "framing": s["framing"], "desc": s["desc"],
                         "prompt": (s["prompt"] + ", " + prj["meta"]["visual_style"]) if s["prompt"] else "",
                         **{k: s[k] for k in ("transition", "style") if s.get(k)}})
            st += s["dur"]
        t += float(b.get("seconds") or 0)
    return {"title": prj["meta"]["title"], "language": prj["meta"]["language"],
            "narration": [{"beat": b["id"], "text": b.get("narration", "")} for b in prj["beats"]],
            "shots": plan}


# ───────────────────────── CLI ─────────────────────────

def main() -> None:
    ap = argparse.ArgumentParser(prog="docos", description="DocOS — system scenopisarski dla dokumentów YouTube")
    sp = ap.add_subparsers(dest="cmd", required=True)
    p = sp.add_parser("new", help="nowy projekt");                p.add_argument("file"); p.add_argument("title")
    p = sp.add_parser("system", help="globalny system prompt");   p.add_argument("file")
    p = sp.add_parser("prompt", help="złota formuła dla beatu");  p.add_argument("file"); p.add_argument("beat")
    p = sp.add_parser("next", help="prompt dla pierwszego nienapisanego beatu"); p.add_argument("file")
    p = sp.add_parser("save", help="zapisz wynik AI do beatu");   p.add_argument("file"); p.add_argument("beat"); p.add_argument("result", help="plik z odpowiedzią lub '-' dla stdin")
    p = sp.add_parser("check", help="kontrola energii, faktów, czasu"); p.add_argument("file")
    p = sp.add_parser("status", help="tabela beatów");            p.add_argument("file")
    p = sp.add_parser("export", help="xlsx + script.md + narration.txt + render_plan.json"); p.add_argument("file"); p.add_argument("--out", default=None)
    p = sp.add_parser("import", help="wczytaj planowanie z xlsx"); p.add_argument("file"); p.add_argument("xlsx")
    a = ap.parse_args()
    path = Path(a.file)

    if a.cmd == "new":
        if path.exists():
            sys.exit(f"{path} już istnieje")
        dump(path, new_project(a.title))
        print(f"Utworzono {path}. Uzupełnij meta, rozdziały, fakty i beaty (albo: export → edytuj xlsx → import).")
        return
    prj = load(path)

    if a.cmd == "system":
        print(system_prompt(prj))
    elif a.cmd == "prompt":
        print(beat_prompt(prj, beat_index(prj, a.beat)))
    elif a.cmd == "next":
        todo = [i for i, b in enumerate(prj["beats"]) if not b.get("narration")]
        if not todo:
            print("Wszystkie beaty napisane. Uruchom: check, potem export.")
            return
        print(f"# NASTĘPNY: {prj['beats'][todo[0]]['id']} (zostało {len(todo)})\n")
        print(beat_prompt(prj, todo[0]))
    elif a.cmd == "save":
        i = beat_index(prj, a.beat)
        text = sys.stdin.read() if a.result == "-" else Path(a.result).read_text(encoding="utf-8")
        r = parse_result(text)
        if not r["narration"]:
            sys.exit("Nie znaleziono sekcji 'Narracja:' — sprawdź format odpowiedzi.")
        prj["beats"][i].update(r, raw=text)
        dump(path, prj)
        print(f"Zapisano {prj['beats'][i]['id']}: {len(r['narration'].split())} słów, {len(r['shots'])} ujęć, "
              f"fakty {r['facts_used'] or '-'}. Połączenie → N-1 dla kolejnego beatu.")
        for cat, msg in check(prj):
            if msg.startswith(prj["beats"][i]["id"] + ":"):
                print(f"  ⚠ [{cat}] {msg}")
    elif a.cmd == "check":
        issues = check(prj)
        for cat, msg in issues:
            print(f"[{cat}] {msg}")
        print("OK — brak problemów." if not issues else f"\n{len(issues)} uwag.")
    elif a.cmd == "status":
        t = 0.0
        print(f"{'ID':6} {'czas':>6} {'segment':14} {'dyn':7} {'E':>2} {'chaos':6} {'słowa':>5} {'ujęć':>4}  turn")
        for b in prj["beats"]:
            d = effective_dynamic(prj, b)
            print(f"{b['id']:6} {int(t // 60):>3}:{int(t % 60):02d} {b.get('segment', ''):14} {d:7} {DYNAMICS[d][0]:>2} "
                  f"{b.get('chaos') or 'NONE':6} {len(b.get('narration', '').split()):>5} {len(b.get('shots', [])):>4}  {b.get('turn', '')}")
            t += float(b.get("seconds") or 0)
        print(f"Razem: {t / 60:.1f} min / cel {prj['meta'].get('target_minutes')} min")
    elif a.cmd == "export":
        out = Path(a.out) if a.out else path.with_suffix("")
        out.mkdir(parents=True, exist_ok=True)
        stem = path.stem
        export_xlsx(prj, out / f"{stem}.xlsx")
        (out / f"{stem}_script.md").write_text(export_script(prj), encoding="utf-8")
        (out / f"{stem}_narration.txt").write_text(
            "\n\n".join(b.get("narration", "") for b in prj["beats"] if b.get("narration")) + "\n", encoding="utf-8")
        (out / f"{stem}_render_plan.json").write_text(json.dumps(render_plan(prj), ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"Wyeksportowano do {out}/ : {stem}.xlsx, {stem}_script.md, {stem}_narration.txt, {stem}_render_plan.json")
    elif a.cmd == "import":
        dump(path, import_xlsx(prj, Path(a.xlsx)))
        print(f"Wczytano planowanie z {a.xlsx} do {path}.")


if __name__ == "__main__":
    main()
