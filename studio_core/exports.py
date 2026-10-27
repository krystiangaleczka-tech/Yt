"""Publication text, measured captions and editable, lossless project workbooks."""
import hashlib
import io
import json

from .models import Project, fps
from .freshness import voice_state
from .voices import alignment_hash


def attribution(project):
    ids={c.asset_id for c in project.sequence.clips if c.asset_id}|{a for c in project.sequence.clips for a in c.asset_ids}
    return [{"id":a.id,"name":a.name,"author":a.author,"source":a.source_url,"license":a.license,"license_url":a.license_url,"attribution":a.attribution} for a in project.assets if a.id in ids and not a.private]


def publication(project, format):
    narration="\n\n".join(s.narration for s in project.scenes if s.narration)
    if format=="txt": return narration
    credits="\n".join(" · ".join(filter(None,[a["name"],a["author"],a["license"],a["source"],a["license_url"],a["attribution"]])) for a in attribution(project))
    facts={fid for scene in project.scenes for fid in scene.fact_ids}|{fid for c in project.sequence.clips for fid in c.fact_ids}
    sources="\n".join(f"- [{f.id}] {f.claim} — {f.source}; {f.url}"+(f"; s. {f.page}" if f.page else "") for f in project.facts if f.id in facts)
    if format=="description": return project.title+"\n\n"+project.story.get("thesis","")+"\n\nŹródła:\n"+sources+"\n\nMateriały i autorzy:\n"+credits
    if format=="credits": return credits
    if format=="md":
        text=f"# {project.title}\n\nRewizja {project.revision} · {project.updated_at}\n\n## Narracja\n\n{narration}\n\n## Ujęcia\n\n"
        for scene in project.scenes:
            text+=f"### {scene.title}\n\n{scene.start/fps(project):.3f}–{(scene.start+scene.duration)/fps(project):.3f} s · {scene.template}\n\n{scene.narration}\n\nFakty: {', '.join(scene.fact_ids)}\n\n"
        return text+"## Źródła\n\n"+sources+"\n\n## Materiały i autorzy\n\n"+credits
    raise ValueError("Nieznany format tekstowy")


def caption_segments(project):
    rate=fps(project); output=[]
    for scene in project.scenes:
        if not scene.narration.strip(): continue
        if voice_state(project,scene)!="current":
            raise ValueError("Przed eksportem napisów odśwież nagranie: "+scene.title)
        if scene.alignment_hash!=alignment_hash(project,scene) or not scene.words:
            raise ValueError("Dopasuj tekst do nagrania lub zatwierdź zmierzone zdania: "+scene.title)
        solo=any(t.solo for t in project.sequence.tracks)
        clips=[c for c in project.sequence.clips if c.scene_id==scene.id and c.asset_id==scene.voice_asset_id and any(t.id==c.track_id and t.role=="narration" and not t.muted and (not solo or t.solo) for t in project.sequence.tracks)]
        if not clips: raise ValueError("Nagranie nie jest słyszalne na osi: "+scene.title)
        for clip in clips:
            current=[]
            for item in scene.words:
                start=clip.start/rate+(float(item["start"])-clip.source_in/rate)/clip.speed
                end=clip.start/rate+(float(item["end"])-clip.source_in/rate)/clip.speed
                if end<=clip.start/rate or start>=(clip.start+clip.duration)/rate: continue
                word={"word":str(item.get("word",item.get("text",""))),"start":max(clip.start/rate,start),"end":min((clip.start+clip.duration)/rate,end)}
                current.append(word)
                if item.get("break_after") or len(current)>=7 or current[-1]["end"]-current[0]["start"]>=3 or word["word"].rstrip().endswith((".","?","!")):
                    output.append({"start":current[0]["start"],"end":current[-1]["end"],"text":" ".join(w["word"].strip() for w in current)}); current=[]
            if current: output.append({"start":current[0]["start"],"end":current[-1]["end"],"text":" ".join(w["word"].strip() for w in current)})
    return sorted(output,key=lambda item:item["start"])


def captions(project, format):
    def stamp(seconds):
        ms=round(seconds*1000); hours,ms=divmod(ms,3600000); minutes,ms=divmod(ms,60000); sec,ms=divmod(ms,1000)
        return f"{hours:02}:{minutes:02}:{sec:02}"+("," if format=="srt" else ".")+f"{ms:03}"
    lines=["WEBVTT\n"] if format=="vtt" else []
    segments=caption_segments(project)
    if not segments: raise ValueError("Brak zmierzonych napisów do eksportu")
    for index,item in enumerate(segments,1):
        lines.append(f"{index}\n{stamp(item['start'])} --> {stamp(item['end'])}\n{item['text']}\n")
    return "\n".join(lines)


def workbook(project):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    book=Workbook(); book.remove(book.active)
    tables={"Projekt":(["id","title","language","channel"],[project.model_dump()]),
            "Fakty":(["id","claim","source","url","page","status","note","source_id","quote"],[f.model_dump() for f in project.facts]),
            "Historia":(["field","text"],[{"field":k,"text":v} for k,v in project.story.items() if k!="status" and not k.startswith("approved_")]),
            "Źródła":(["id","title","author","edition","identifier","url"],[s.model_dump() for s in project.sources]),
            "Materiały":(["id","name","author","license","license_url","attribution","source_url"],[a.model_dump() for a in project.assets]),
            "Fragmenty":(["id","title","chapter","goal","narration","turn","energy","dynamics","typology","chaos","fact_ids"],[b.model_dump() for b in project.beats]),
            "Ujęcia":(["id","title","narration","start","duration","beat_id","fact_ids","template","parameters"],[s.model_dump() for s in project.scenes]),
            "Klipy":(["id","track_id","asset_id","start","duration","source_in","speed","volume","text","transition","transition_duration"],[c.model_dump() for c in project.sequence.clips])}
    for name,(columns,rows) in tables.items():
        sheet=book.create_sheet(name); sheet.append(columns)
        for row in rows:
            sheet.append([json.dumps(row.get(k),ensure_ascii=False) if isinstance(row.get(k),(list,dict)) else row.get(k) for k in columns])
        sheet.freeze_panes="A2"; sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="30364A")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width=min(70,max(15,len(str(column[0].value))+4))
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment=Alignment(vertical="top",wrap_text=True)
                if isinstance(cell.value,str): cell.data_type="s"
    manifest=book.create_sheet("_manifest"); manifest.sheet_state="veryHidden"
    raw=json.dumps(project.model_dump(),ensure_ascii=False)
    manifest.append(["revision",project.revision]); manifest.append(["sha256",hashlib.sha256(raw.encode()).hexdigest()])
    for offset in range(0,len(raw),30000): manifest.append(["json",raw[offset:offset+30000]])
    for row in manifest:
        for cell in row:
            if isinstance(cell.value,str):cell.data_type="s"
    output=io.BytesIO(); book.save(output); return output.getvalue()


def read_workbook(content):
    from openpyxl import load_workbook
    import zipfile
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        if sum(item.file_size for item in archive.infolist())>64*1024*1024:raise ValueError("Arkusz przekracza limit 64 MB po rozpakowaniu")
    book=load_workbook(io.BytesIO(content),data_only=False)
    if any(cell.data_type=="f" for sheet in book for row in sheet for cell in row):
        raise ValueError("Formuły nie są danymi projektu")
    if "_manifest" not in book: raise ValueError("To nie jest arkusz DocOS Studio")
    rows=list(book["_manifest"].values); raw="".join(str(row[1]) for row in rows if row[0]=="json")
    if hashlib.sha256(raw.encode()).hexdigest()!=rows[1][1]: raise ValueError("Manifest XLSX jest uszkodzony")
    project=Project.model_validate(json.loads(raw)); data=project.model_dump()
    for name,field in (("Fakty","facts"),("Źródła","sources"),("Materiały","assets"),("Fragmenty","beats"),("Ujęcia","scenes"),("Klipy",None)):
        if name not in book: continue
        values=list(book[name].values); columns=list(values[0]); originals={item.id:item.model_dump() for item in (getattr(project,field) if field else project.sequence.clips)}
        updated=[]
        for values_row in values[1:]:
            if not values_row[0]: continue
            entry=originals.get(str(values_row[0]))
            if entry is None: raise ValueError("Nie zmieniaj ID w arkuszu. Nowe elementy dodaj w aplikacji.")
            for key,value in zip(columns,values_row):
                if isinstance(entry.get(key),(list,dict)): value=json.loads(value or ("[]" if isinstance(entry[key],list) else "{}"))
                entry[key]=value if value is not None else ("" if isinstance(entry.get(key),str) else entry.get(key))
            updated.append(entry)
        if field: data[field]=updated
        else: data["sequence"]["clips"]=updated
    if "Projekt" in book:
        values=list(book["Projekt"].values)
        if len(values)>1:
            for key,value in zip(values[0],values[1]):
                if key in ("title","language","channel") and value is not None: data[key]=str(value)
    if "Historia" in book:
        for key,value in list(book["Historia"].values)[1:]:
            if key in data["story"] and key!="status" and not str(key).startswith("approved_"):data["story"][key]=str(value or "")
    return Project.model_validate(data)
